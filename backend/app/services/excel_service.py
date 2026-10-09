import io
import openpyxl
from typing import List, Dict, Any


class ExcelService:
    @staticmethod
    def parse_student_excel(file_bytes: bytes) -> List[Dict[str, Any]]:
        wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active

        students = []
        # Expected header row 1: [Name, Email, Student ID, NISN, Gender (M/F), Phone, Address]
        for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_idx == 1:
                continue  # skip header
            if not row or not any(row):
                continue

            name = str(row[0]).strip() if len(row) > 0 and row[0] is not None else None
            email = str(row[1]).strip() if len(row) > 1 and row[1] is not None else None
            nis = str(row[2]).strip() if len(row) > 2 and row[2] is not None else None
            nisn = str(row[3]).strip() if len(row) > 3 and row[3] is not None else None
            gender = str(row[4]).strip().upper() if len(row) > 4 and row[4] is not None else "L"
            phone = str(row[5]).strip() if len(row) > 5 and row[5] is not None else None
            address = str(row[6]).strip() if len(row) > 6 and row[6] is not None else None

            if name and email and nis:
                students.append({
                    "name": name,
                    "email": email,
                    "nis": nis,
                    "nisn": nisn,
                    "gender": gender if gender in ["L", "P", "M", "F"] else "L",
                    "phone": phone,
                    "address": address,
                })

        return students

    @staticmethod
    def parse_questions_excel(file_bytes: bytes) -> List[Dict[str, Any]]:
        wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active

        questions = []
        # Expected header: [Question Text, Type, Option A, Option B, Option C, Option D, Correct Answer, Points, Difficulty]
        for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_idx == 1:
                continue
            if not row or not any(row):
                continue

            q_text = str(row[0]).strip() if len(row) > 0 and row[0] is not None else None
            q_type = str(row[1]).strip().lower() if len(row) > 1 and row[1] is not None else "multiple_choice"
            opt_a = str(row[2]).strip() if len(row) > 2 and row[2] is not None else None
            opt_b = str(row[3]).strip() if len(row) > 3 and row[3] is not None else None
            opt_c = str(row[4]).strip() if len(row) > 4 and row[4] is not None else None
            opt_d = str(row[5]).strip() if len(row) > 5 and row[5] is not None else None
            correct = str(row[6]).strip().upper() if len(row) > 6 and row[6] is not None else "A"
            points = int(row[7]) if len(row) > 7 and row[7] is not None and str(row[7]).isdigit() else 1
            diff = str(row[8]).strip().lower() if len(row) > 8 and row[8] is not None else "medium"

            options = []
            if opt_a:
                options.append({"key": "A", "text": opt_a})
            if opt_b:
                options.append({"key": "B", "text": opt_b})
            if opt_c:
                options.append({"key": "C", "text": opt_c})
            if opt_d:
                options.append({"key": "D", "text": opt_d})

            if q_text:
                questions.append({
                    "question_text": q_text,
                    "question_type": q_type if q_type in ["multiple_choice", "essay", "short_answer", "true_false"] else "multiple_choice",
                    "options": options if options else None,
                    "correct_answer": correct,
                    "points": points,
                    "difficulty": diff if diff in ["easy", "medium", "hard"] else "medium",
                })

        return questions


excel_service = ExcelService()
